#!/usr/bin/env python3
"""Perfila los datos reales de un modelo de Odoo y clasifica el origen de cada campo.

Reporte de comprensión: contra CUALQUIER instancia alcanzable por el MCP odoo-mcp-19,
responde empíricamente "¿qué campos de este modelo tienen datos y cuáles migran?".
Odoo no guarda proveniencia por registro, pero ``ir.model.fields`` revela cómo se
puebla cada campo, y eso decide el alcance de una migración:

    CUSTOM           -> state='manual'  (x_studio_*, campo agregado a mano) ......... migra
    SYSTEM: related  -> related != False (espejo de otro registro) .................. no migra
    SYSTEM: computed -> compute o store=False (lo calcula Odoo) ..................... no migra
    SYSTEM: audit    -> create_/write_/id (lo pone el ORM) ......................... no migra
    SYSTEM: readonly -> readonly=True (lo setea la lógica de negocio) .............. no migra
    USER             -> almacenado + escribible + no calculado ..................... migra

Salidas (en el directorio actual):
    <out>.json   : conteo de registros, campos poblados, meta y origen de cada uno.
    <out>.xlsx   : 3 hojas — "Resumen campos" (todos + origen + cobertura %, migrables
                   en verde), "Datos (todos)" y "Migracion (USUARIO)". Requiere openpyxl.

No depende de ningún proyecto: descubre el wrapper del server MCP leyendo el
``.mcp.json`` del directorio actual (o de un ancestro), así apunta a la instancia
que ese proyecto tenga configurada. Ver ``--wrapper`` para forzarlo.

Uso:
    python report.py                                   # res.partner, todos los registros
    python report.py --model res.partner --domain '[["supplier_rank",">",0]]'
    python report.py --model sale.order --out ventas   # otro modelo
    python report.py --include-archived                # incluye registros archivados
"""
from __future__ import annotations

import argparse
import json
import subprocess
import sys
import time
from collections import Counter
from pathlib import Path
from typing import Any

MAX_PAGE = 1000  # tope de registros por search_read que impone el server MCP
# Valores que cuentan como "vacío" para decidir si un campo tiene datos.
# Nota deliberada: 0 y 0.0 se tratan como vacíos (perfil de cobertura, no de valores);
# un 0 legítimo — p.ej. una cantidad en cero — se contaría como sin datos.
EMPTY = (False, None, "", [], {}, 0, 0.0)
AUDIT = {"create_uid", "create_date", "write_uid", "write_date", "id"}
MIGRATABLE = {"USER", "CUSTOM"}


def log(*a: object) -> None:
    print(*a, file=sys.stderr, flush=True)


# --------------------------------------------------------------------------- #
#  Descubrimiento del wrapper del MCP                                          #
# --------------------------------------------------------------------------- #

def discover_wrapper(explicit: str | None) -> list[str]:
    """Devuelve el argv para lanzar el server MCP (stdio).

    Orden: ``--wrapper`` explícito → variable ``ODOO_MCP_WRAPPER`` → el ``command``
    del ``.mcp.json`` (o ``.mcp.local.json``) del cwd o de un ancestro. Prefiere la
    clave de server cuyo nombre contenga "odoo"; si no, la primera cuyo command
    termine en ``run-docker.sh``.
    """
    if explicit:
        return [explicit]
    import os
    env = os.environ.get("ODOO_MCP_WRAPPER")
    if env:
        return [env]

    for base in [Path.cwd(), *Path.cwd().parents]:
        for fname in (".mcp.json", ".mcp.local.json"):
            cfg = base / fname
            if not cfg.is_file():
                continue
            try:
                servers = json.loads(cfg.read_text()).get("mcpServers", {})
            except (json.JSONDecodeError, OSError):
                continue
            # candidato preferido: nombre con "odoo"; luego cualquiera run-docker.sh
            ordered = sorted(servers.items(), key=lambda kv: "odoo" not in kv[0].lower())
            for _name, spec in ordered:
                cmd = spec.get("command")
                if not cmd:
                    continue
                argv = [cmd] + list(spec.get("args", []))
                if "odoo" in _name.lower() or str(cmd).endswith("run-docker.sh"):
                    log(f"[wrapper] {cfg} :: {_name} -> {' '.join(argv)}")
                    return argv
    raise SystemExit(
        "No encontré cómo lanzar el MCP. Pasá --wrapper /ruta/a/run-docker.sh, "
        "definí ODOO_MCP_WRAPPER, o corré esto dentro de un proyecto con .mcp.json."
    )


# --------------------------------------------------------------------------- #
#  Cliente MCP stdio mínimo y autocontenido                                    #
# --------------------------------------------------------------------------- #

class MCP:
    """Cliente stdio mínimo para odoo-mcp-19. stdin queda abierto hasta recibir
    cada respuesta: FastMCP trata el EOF como shutdown y aborta la llamada en vuelo."""

    def __init__(self, argv: list[str], timeout: int = 180) -> None:
        self.timeout = timeout
        self.p = subprocess.Popen(
            argv, stdin=subprocess.PIPE, stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL, text=True, bufsize=1,
        )
        self._id = 0
        self._send({"method": "initialize", "params": {
            "protocolVersion": "2024-11-05", "capabilities": {},
            "clientInfo": {"name": "odoo-data-report", "version": "1"}}})
        self._await(self._id)
        self._notify("notifications/initialized")

    def _send(self, msg: dict[str, Any]) -> int:
        self._id += 1
        self.p.stdin.write(json.dumps({"jsonrpc": "2.0", "id": self._id, **msg}) + "\n")
        self.p.stdin.flush()
        return self._id

    def _notify(self, method: str) -> None:
        self.p.stdin.write(json.dumps({"jsonrpc": "2.0", "method": method}) + "\n")
        self.p.stdin.flush()

    def _await(self, rid: int) -> dict[str, Any]:
        start = time.time()
        while time.time() - start < self.timeout:
            line = self.p.stdout.readline()
            if not line:
                raise RuntimeError("el server MCP cerró el stream (¿Docker corriendo?)")
            try:
                d = json.loads(line)
            except json.JSONDecodeError:
                continue
            if d.get("id") == rid:
                return d
        raise TimeoutError(f"sin respuesta del MCP para id={rid} en {self.timeout}s")

    def execute_method(self, model: str, method: str,
                       args: list[Any] | None = None, **kwargs: Any) -> Any:
        arguments: dict[str, Any] = {"model": model, "method": method}
        if args is not None:
            arguments["args_json"] = json.dumps(args)
        if kwargs:
            arguments["kwargs_json"] = json.dumps(kwargs)
        rid = self._send({"method": "tools/call",
                          "params": {"name": "execute_method", "arguments": arguments}})
        resp = self._await(rid)
        if "error" in resp:
            raise RuntimeError(resp["error"])
        payload = json.loads(resp["result"]["content"][0]["text"])
        if not payload.get("success"):
            raise RuntimeError(payload.get("error") or payload)
        return payload["result"]

    def search_count(self, model: str, domain: list, ctx: dict | None = None) -> int:
        kw = {"context": ctx} if ctx else {}
        return int(self.execute_method(model, "search_count", args=[domain], **kw))

    def fetch_all(self, model: str, domain: list, ctx: dict | None = None) -> list[dict]:
        total = self.search_count(model, domain, ctx)
        rows: list[dict] = []
        kw: dict[str, Any] = {"domain": domain, "limit": MAX_PAGE}
        if ctx:
            kw["context"] = ctx
        while len(rows) < total:
            page = self.execute_method(model, "search_read", offset=len(rows), **kw)
            if not page:
                break
            rows.extend(page)
            log(f"  ... {len(rows)}/{total}")
        return rows

    def close(self) -> None:
        try:
            self.p.stdin.close()
        finally:
            self.p.terminate()


# --------------------------------------------------------------------------- #
#  Clasificación de origen                                                     #
# --------------------------------------------------------------------------- #

def origin(name: str, meta: dict[str, dict]) -> str:
    m = meta.get(name, {})
    if m.get("state") == "manual":
        return "CUSTOM"
    if m.get("related"):
        return "SYSTEM: related"
    if m.get("compute") or not m.get("store"):
        return "SYSTEM: computed"
    if name in AUDIT:
        return "SYSTEM: audit"
    if m.get("readonly"):
        return "SYSTEM: readonly"
    return "USER"


def cell(v: Any) -> Any:
    """Aplana un valor de Odoo para una celda de planilla."""
    if isinstance(v, (list, tuple)):
        if len(v) == 2 and isinstance(v[0], int) and isinstance(v[1], str):
            return v[1]                       # many2one -> "nombre"
        return ", ".join(str(x) for x in v)   # x2many -> ids
    if isinstance(v, dict):
        return json.dumps(v, default=str)
    return v


# --------------------------------------------------------------------------- #
#  Export XLSX                                                                 #
# --------------------------------------------------------------------------- #

def write_xlsx(path: Path, records: list[dict], all_fields: list[str],
               populated: list[str], labels: dict, counts: dict, origins: dict) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        log("[xlsx] openpyxl no está instalado — omito el .xlsx (el .json sí se generó). "
            "Instalá con: pip install openpyxl")
        return False

    n = len(records) or 1
    wb = Workbook()
    green = PatternFill("solid", fgColor="C6EFCE")

    ws = wb.active
    ws.title = "Resumen campos"
    ws.append(["campo", "etiqueta", "tipo", "origen", "registros_con_dato", "pct"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for f in sorted(all_fields, key=lambda x: (x not in populated,
                                               origins.get(x, "") not in MIGRATABLE,
                                               -counts[x])):
        o = origins.get(f, "SYSTEM: vacío")
        lm = labels.get(f, {})
        ws.append([f, lm.get("string", ""), lm.get("type", ""), o,
                   counts[f], round(100 * counts[f] / n, 1)])
        if o in MIGRATABLE:
            for c in ws[ws.max_row]:
                c.fill = green

    data_fields = [f for f in populated
                   if labels.get(f, {}).get("type") != "binary"
                   and not f.startswith(("image_", "avatar_"))]
    ws2 = wb.create_sheet("Datos (todos)")
    ws2.append(data_fields)
    for c in ws2[1]:
        c.font = Font(bold=True)
    for r in records:
        ws2.append([cell(r.get(f)) for f in data_fields])

    migrate_fields = [f for f in data_fields if origins.get(f) in MIGRATABLE]
    ws3 = wb.create_sheet("Migracion (USUARIO)")
    ws3.append(migrate_fields)
    for c in ws3[1]:
        c.font = Font(bold=True)
    for r in records:
        ws3.append([cell(r.get(f)) for f in migrate_fields])

    for sh in (ws, ws2, ws3):
        sh.freeze_panes = "A2"
    for i, w in enumerate([34, 36, 12, 18, 20, 8], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)
    return True


# --------------------------------------------------------------------------- #
#  Main                                                                        #
# --------------------------------------------------------------------------- #

def main() -> None:
    ap = argparse.ArgumentParser(description="Reporte de comprensión de un modelo de Odoo.")
    ap.add_argument("--model", default="res.partner", help="modelo (default: res.partner)")
    ap.add_argument("--domain", default="[]", help='dominio JSON (default: [] = todos)')
    ap.add_argument("--out", default=None, help="basename de salida (default: <modelo>_report)")
    ap.add_argument("--wrapper", default=None, help="ruta a run-docker.sh (default: auto)")
    ap.add_argument("--include-archived", action="store_true", help="incluir registros archivados")
    ap.add_argument("--no-xlsx", action="store_true", help="solo JSON, sin planilla")
    args = ap.parse_args()

    model = args.model
    domain = json.loads(args.domain)
    out = args.out or model.replace(".", "_") + "_report"
    ctx = {"active_test": False} if args.include_archived else None

    mcp = MCP(discover_wrapper(args.wrapper))
    try:
        total = mcp.search_count(model, domain, ctx)
        log(f"{model}: {total} registros para el dominio {domain}"
            + (" (incluye archivados)" if ctx else ""))
        if total == 0:
            raise SystemExit("El dominio no devolvió registros; nada para perfilar.")
        records = mcp.fetch_all(model, domain, ctx)

        all_fields = sorted({f for r in records for f in r})
        populated = sorted({f for r in records for f, v in r.items() if v not in EMPTY})
        counts = {f: sum(1 for r in records if r.get(f) not in EMPTY) for f in all_fields}

        labels = mcp.execute_method(model, "fields_get", args=[populated],
                                    attributes=["string", "type"])
        meta_rows = mcp.execute_method(
            "ir.model.fields", "search_read",
            domain=[["model", "=", model], ["name", "in", populated]],
            fields=["name", "store", "readonly", "compute", "related", "state"],
            limit=0)  # limit=0 evita el cap por defecto de ~100 filas
        meta = {m["name"]: m for m in meta_rows}
        origins = {f: origin(f, meta) for f in populated}
    finally:
        pass  # mantené el MCP abierto hasta acá; se cierra abajo

    # --- persistir JSON ---
    summary = {
        "model": model, "domain": domain, "include_archived": bool(ctx),
        "record_count": len(records),
        "field_count": len(all_fields), "populated_count": len(populated),
        "fields": {f: {"label": labels.get(f, {}).get("string", ""),
                       "type": labels.get(f, {}).get("type", ""),
                       "origin": origins[f],
                       "records_with_data": counts[f],
                       "pct": round(100 * counts[f] / (len(records) or 1), 1)}
                   for f in populated},
    }
    Path(f"{out}.json").write_text(json.dumps(summary, indent=2, default=str))

    wrote_xlsx = False
    if not args.no_xlsx:
        wrote_xlsx = write_xlsx(Path(f"{out}.xlsx"), records, all_fields,
                                populated, labels, counts, origins)
    mcp.close()

    # --- resumen didáctico a stdout ---
    by_origin = Counter(origins[f] for f in populated)
    migrables = sorted((f for f in populated if origins[f] in MIGRATABLE),
                       key=lambda f: -counts[f])
    print(f"\n=== {model} · {len(records)} registros ===")
    print(f"{len(populated)} de {len(all_fields)} campos devueltos tienen datos.\n")
    print("Campos poblados por origen:")
    for o, k in by_origin.most_common():
        print(f"  {o:<18} {k}")
    print(f"\nMigrables (USER/CUSTOM) con datos: {len(migrables)}")
    for f in migrables:
        lm = labels.get(f, {})
        print(f"  {f:<34} {counts[f]:>6} reg ({summary['fields'][f]['pct']:>5}%)  "
              f"{lm.get('string', '')}")
    print(f"\nGenerado: {out}.json" + (f" + {out}.xlsx" if wrote_xlsx else ""))


if __name__ == "__main__":
    main()
